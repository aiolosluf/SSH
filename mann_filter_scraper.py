import json
import os
import re
import time
import urllib.request
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

# ======== 配置区（按你的实际路径修改）========
EXCEL_PATH = r"C:\Users\admin\OneDrive\桌面\python\web crawler\MANN_PRODUCT_LIST.xlsx"
OUTPUT_EXCEL_PATH = EXCEL_PATH
IMAGE_FOLDER = r"C:\Users\admin\OneDrive\桌面\python\MANN"

# A列可以放 part no（例如 CUK23005-2）或完整产品URL
INPUT_COL = 1
# 输出列：可按需改
COL_IMAGE = 2
COL_DESC = 3
COL_DIMENSIONS = 4
COL_GTIN = 5
COL_STATUS = 6

BASE_URL = "https://www.mann-filter.com/au-en/catalog/search-results/product.html"


def setup_driver():
    options = Options()
    # 如果你想看浏览器过程，就注释掉下一行
    # options.add_argument("--headless=new")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--disable-blink-features=AutomationControlled")
    driver = webdriver.Chrome(options=options)
    return driver


def ensure_folder(path):
    if not os.path.exists(path):
        os.makedirs(path)


def normalize_partno(partno):
    return re.sub(r"\s+", "", str(partno).strip())


def build_product_url(raw_value):
    """
    支持两种输入：
    1) 已经是完整URL
    2) 只是 part no，例如 CUK23005-2
    """
    value = str(raw_value).strip()
    if value.lower().startswith("http"):
        return value

    partno = normalize_partno(value)
    slug = partno.lower()
    return f"{BASE_URL}/{slug}_mann-filter.html"


def safe_filename(name):
    return re.sub(r"[\\/:*?\"<>|]", "_", name)


def get_jsonld_blocks(soup):
    blocks = []
    for tag in soup.find_all("script", attrs={"type": "application/ld+json"}):
        txt = (tag.string or tag.get_text() or "").strip()
        if not txt:
            continue
        try:
            data = json.loads(txt)
            blocks.append(data)
        except json.JSONDecodeError:
            # 有些站点会塞多个JSON对象或注释，简单跳过
            continue
    return blocks


def flatten_jsonld(data):
    """把 json-ld 拍平，方便找 Product 节点"""
    result = []
    if isinstance(data, dict):
        result.append(data)
        for v in data.values():
            result.extend(flatten_jsonld(v))
    elif isinstance(data, list):
        for item in data:
            result.extend(flatten_jsonld(item))
    return result


def extract_from_jsonld(soup):
    product = {}
    blocks = get_jsonld_blocks(soup)
    nodes = []
    for b in blocks:
        nodes.extend(flatten_jsonld(b))

    product_nodes = []
    for n in nodes:
        if not isinstance(n, dict):
            continue
        type_field = n.get("@type", "")
        if isinstance(type_field, list):
            type_text = " ".join(type_field).lower()
        else:
            type_text = str(type_field).lower()
        if "product" in type_text:
            product_nodes.append(n)

    if product_nodes:
        p = product_nodes[0]
        product["description"] = p.get("description")
        product["gtin"] = p.get("gtin") or p.get("gtin13") or p.get("gtin12")
        img = p.get("image")
        if isinstance(img, list):
            product["image"] = img[0] if img else None
        else:
            product["image"] = img

    return product


def find_text_by_label(soup, labels):
    """按标签名查值，比如 GTIN、Dimension。"""
    label_regex = re.compile("|".join([re.escape(x) for x in labels]), re.I)

    # 常见 key-value 结构
    for el in soup.find_all(string=label_regex):
        key_text = el.strip().lower()
        if not key_text:
            continue

        parent = el.parent
        if not parent:
            continue

        # 尝试邻近元素取值
        candidates = []
        if parent.find_next_sibling():
            candidates.append(parent.find_next_sibling())
        if parent.parent and parent.parent.find_next_sibling():
            candidates.append(parent.parent.find_next_sibling())

        for c in candidates:
            text = c.get_text(" ", strip=True)
            if text:
                return text

    return None


def collect_dimensions_text(soup):
    """尽量抓取尺寸相关字段，拼成一段文本。"""
    patterns = [
        r"dimension", r"height", r"width", r"length", r"diameter", r"outer", r"inner", r"size"
    ]
    reg = re.compile("|".join(patterns), re.I)

    lines = []

    # 1) 表格行
    for tr in soup.select("tr"):
        tds = tr.find_all(["th", "td"])
        if len(tds) >= 2:
            key = tds[0].get_text(" ", strip=True)
            val = tds[1].get_text(" ", strip=True)
            if reg.search(key) and val:
                lines.append(f"{key}: {val}")

    # 2) 列表结构
    for li in soup.select("li"):
        txt = li.get_text(" ", strip=True)
        if reg.search(txt):
            lines.append(txt)

    # 去重保序
    dedup = []
    seen = set()
    for x in lines:
        if x not in seen:
            seen.add(x)
            dedup.append(x)

    return " | ".join(dedup) if dedup else None


def download_image(image_url, save_folder, partno):
    if not image_url:
        return None

    ensure_folder(save_folder)

    path_obj = urlparse(image_url)
    ext = os.path.splitext(path_obj.path)[1].lower()
    if ext not in [".jpg", ".jpeg", ".png", ".webp"]:
        ext = ".jpg"

    filename = safe_filename(partno) + ext
    filepath = os.path.join(save_folder, filename)

    headers = {
        "User-Agent": "Mozilla/5.0",
        "Referer": "https://www.mann-filter.com/",
    }
    req = urllib.request.Request(image_url, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as resp, open(filepath, "wb") as f:
        f.write(resp.read())

    return filepath


def scrape_product(driver, url, partno):
    result = {
        "image_url": None,
        "description": None,
        "dimensions": None,
        "gtin": None,
        "status": "",
    }

    driver.get(url)

    try:
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.TAG_NAME, "body"))
        )
    except TimeoutException:
        result["status"] = "timeout"
        return result

    time.sleep(1.5)
    soup = BeautifulSoup(driver.page_source, "html.parser")

    # 1) 先从 json-ld 抓（通常最准）
    j = extract_from_jsonld(soup)

    description = j.get("description")
    gtin = j.get("gtin")
    image_url = j.get("image")

    # 2) 页面兜底
    if not description:
        meta_desc = soup.find("meta", attrs={"name": "description"})
        if meta_desc and meta_desc.get("content"):
            description = meta_desc["content"].strip()

    if not gtin:
        # 文本模式提取 GTIN（8~14位）
        page_text = soup.get_text(" ", strip=True)
        m = re.search(r"GTIN\s*[:#]?\s*(\d{8,14})", page_text, re.I)
        if m:
            gtin = m.group(1)
        else:
            gtin = find_text_by_label(soup, ["GTIN", "EAN"])

    if not image_url:
        og_img = soup.find("meta", attrs={"property": "og:image"})
        if og_img and og_img.get("content"):
            image_url = og_img["content"].strip()

    dimensions = collect_dimensions_text(soup)
    if not dimensions:
        dimensions = find_text_by_label(soup, ["Dimensions", "Size", "Height", "Width", "Length", "Diameter"])

    result["image_url"] = image_url
    result["description"] = description
    result["dimensions"] = dimensions
    result["gtin"] = gtin
    result["status"] = "ok"

    return result


def main():
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    ensure_folder(IMAGE_FOLDER)

    driver = setup_driver()

    try:
        for row in range(2, ws.max_row + 1):
            raw_input = ws.cell(row=row, column=INPUT_COL).value
            if not raw_input:
                continue

            partno = normalize_partno(raw_input)
            url = build_product_url(raw_input)
            print(f"[{row}] 抓取: {partno} -> {url}")

            try:
                data = scrape_product(driver, url, partno)

                image_local_path = None
                if data["image_url"]:
                    try:
                        image_local_path = download_image(data["image_url"], IMAGE_FOLDER, partno)
                    except Exception as e:
                        print(f"  图片下载失败: {e}")

                ws.cell(row=row, column=COL_IMAGE).value = image_local_path or data["image_url"]
                ws.cell(row=row, column=COL_DESC).value = data["description"]
                ws.cell(row=row, column=COL_DIMENSIONS).value = data["dimensions"]
                ws.cell(row=row, column=COL_GTIN).value = data["gtin"]
                ws.cell(row=row, column=COL_STATUS).value = data["status"]

            except Exception as e:
                ws.cell(row=row, column=COL_STATUS).value = f"error: {e}"
                print(f"  抓取失败: {e}")

            wb.save(OUTPUT_EXCEL_PATH)

    finally:
        driver.quit()
        wb.save(OUTPUT_EXCEL_PATH)

    print("全部完成")


if __name__ == "__main__":
    main()
